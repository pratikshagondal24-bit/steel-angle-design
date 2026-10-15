from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\siste\Documents\Codex\2026-10-13\i-x20")
OUT = ROOT / "FA2_Activity_Report_Steel_Angle_Design.docx"
SCREEN = Path(r"C:\Users\siste\OneDrive\Pictures\Screenshots")
PHOTOS = Path(r"C:\Users\siste\Downloads\Telegram Desktop")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = Inches(.68)
sec.bottom_margin = Inches(.68)
sec.left_margin = Inches(.78)
sec.right_margin = Inches(.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.2)
styles["Normal"].font.color.rgb = RGBColor(45, 55, 65)
styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size, color in [("Title", 27, "17324D"), ("Heading 1", 19, "17324D"), ("Heading 2", 13, "126782"), ("Heading 3", 11, "126782")]:
    s = styles[name]
    s.font.name = "Aptos Display"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)

header = sec.header.paragraphs[0]
header.text = "FA ACTIVITY  |  DESIGN OF STEEL STRUCTURES"
header.style = "Caption"
header.runs[0].font.color.rgb = RGBColor(90, 110, 125)
header.runs[0].font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Educational prototype  •  IS 800:2007  |  Page ")
run.font.size = Pt(8)
fld = footer.add_run()
fld._r.append(__import__("docx").oxml.OxmlElement("w:fldChar"))
fld._r[-1].set(__import__("docx").oxml.ns.qn("w:fldCharType"), "begin")
instr = __import__("docx").oxml.OxmlElement("w:instrText")
instr.set(__import__("docx").oxml.ns.qn("xml:space"), "preserve")
instr.text = " PAGE "
fld._r.append(instr)
end = __import__("docx").oxml.OxmlElement("w:fldChar")
end.set(__import__("docx").oxml.ns.qn("w:fldCharType"), "end")
fld._r.append(end)


def para(text="", style=None, bold_lead=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def bullet(text):
    return para(text, "List Bullet")


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Shading Accent 1"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def figure(path, caption, width=6.55):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(caption, style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)


def callout(text):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Light Shading Accent 1"
    cell = t.cell(0, 0)
    cell.text = text
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(5)
        for r in p.runs:
            r.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(56)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("FA ACTIVITY REPORT")
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(18, 103, 130)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
p.add_run("AI Application for Design of Steel Structural Members as per IS 800:2007")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
r = p.add_run("Steel angle tension member checking with Python and Streamlit")
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(80, 98, 112)

doc.add_paragraph(" ")
table(["Student details", "Information"], [
    ("Name", "________________________________________"),
    ("Class / Division", "T.Y. Civil Engineering / __________"),
    ("Roll number", "________________________________________"),
    ("Subject", "Design of Steel Structures"),
    ("Faculty / Institute", "________________________________________"),
    ("Academic year", "________________________________________"),
], [2.0, 4.8])
doc.add_paragraph(" ")
callout("Activity summary: A beginner-friendly Streamlit application was developed to check steel angle tension members for gross-section yielding, net-section rupture, and block shear. The calculation engine is deterministic and displays formulas, intermediate values, design strengths, utilization, and pass/fail status. The submitted examples were also compared with handwritten calculations; differences are recorded for verification rather than presented as confirmed agreement.")
para("Submitted as part of the FA activity in Design of Steel Structures", align=WD_ALIGN_PARAGRAPH.CENTER)

# Page 2: purpose and design
doc.add_page_break()
doc.add_heading("1  Project overview", level=1)
para("This activity applies basic programming to a structural design workflow. The user selects a tensile member, enters the factored axial force and material strengths, enters the bolt connection geometry, and finally enters an angle designation. The application looks up the angle gross area in the supplied SP 6 (Part 1):1964 catalogue transcription and calculates the three included tension limit states.")
doc.add_heading("Aim", level=2)
para("To create an educational web application that performs transparent preliminary checks of a selected steel angle tension member using calculation procedures associated with IS 800:2007.")
doc.add_heading("Objectives", level=2)
for text in [
    "Collect design inputs through a guided, four-step Streamlit interface.",
    "Recognize supported equal and unequal angle designations from the steel-book catalogue.",
    "Calculate yielding, angle rupture, and block shear strength with stated assumptions and partial safety factors.",
    "Show each check separately and compare the design strength with the applied factored force.",
    "Explain results in simple language without using a language model to generate engineering calculations.",
    "Compare example outputs with hand calculations and document any unresolved differences.",
]: bullet(text)
doc.add_heading("Application structure", level=2)
table(["File / component", "Role"], [
    ("app.py", "Streamlit interface, guided inputs, result display and educational explanation."),
    ("design.py", "Input validation and deterministic engineering calculations."),
    ("angle_sections.py", "Angle designation parser and section area catalogue."),
    ("test_design.py", "Automated checks for catalogue lookup, formulas, double angles and invalid inputs."),
    ("requirements.txt / README.md", "Dependencies and instructions for installation and use."),
], [2.0, 4.8])
doc.add_heading("User workflow", level=2)
para("Member type → factored force and material strengths → bolt details and connected leg → angle designation → calculation results. The section is requested last, as required for the guided activity.")
figure(SCREEN / "Screenshot 2026-10-14 183305.png", "Figure 1. Application landing screen and member-type selection.", 6.5)

# Page 3: design basis
doc.add_page_break()
doc.add_heading("2  Design basis and calculation method", level=1)
para("The prototype implements three tension-member checks identified in IS 800:2007: gross-section yielding (Clause 6.2), rupture of a single angle connected by one leg (Clause 6.3.3), and block shear at a bolted connection (Clause 6.4.1). Clause numbering and the detailed equations, limits, definitions, and amendments must be checked against the applicable licensed copy of the standard before academic submission or any real design use.")
table(["Check", "Expression used by the program", "Interpretation"], [
    ("Gross yielding", "Tdg = Ag fy / γm0", "Gross area resistance in tension; γm0 = 1.10 in this prototype."),
    ("Single-angle rupture", "Tdn = 0.9 Anc fu / γm1 + β Ago fy / γm0", "Connected-leg net area and outstanding-leg gross area; β represents shear-lag effects. γm1 = 1.25."),
    ("Block shear", "Tdb = min(Avg fy/(√3 γm0) + 0.9 Atn fu/γm1, 0.9 Avn fu/(√3 γm1) + Atg fy/γm0)", "Two block-shear failure combinations are calculated; the smaller resistance governs."),
], [1.15, 3.0, 2.65])
doc.add_heading("Decision rule", level=2)
para("For each check, the program reports Satisfies when the applied factored force Tu is no greater than the calculated strength for that limit state. Overall design strength is the minimum of the three strengths. The displayed utilization is Tu/Td; a value at or below 1.0 passes the included checks.")
doc.add_heading("Units and material values", level=2)
para("Dimensions are entered in mm, areas are reported in mm², stresses in MPa (N/mm²), and forces in kN. The user supplies fy and fu; the app does not infer values from a steel grade. Example inputs use fy = 250 MPa and fu = 410 MPa, which must be verified for the selected product, grade, and thickness range.")
doc.add_heading("Important modelling assumptions", level=2)
for text in [
    "The catalogue areas are transcribed from the supplied historical SP 6 (Part 1):1964, Tables III and IV; verify designation and area against the source steel book.",
    "The connection model uses one bolt line in one leg. The example starts with five bolts; pitch, end distance, nominal bolt diameter, hole diameter, connected leg, and transverse block path can be edited.",
    "Heel-to-bolt-line gauge is omitted from the form and inferred from connected-leg width minus Lt. This geometrical assumption must match the connection drawing.",
    "A double angle assumes one angle on each side of the gusset with equal force sharing. The app does not design bolt strength, bearing, gusset, eccentricity, compression buckling, or other omitted limit states.",
]: bullet(text)

# Page 4: case A
doc.add_page_break()
doc.add_heading("3  Example one with a double angle", level=1)
para("The first submitted example uses a symmetric double angle to carry a factored axial force of 375 kN. The app screenshot records the following entered material and section data.")
table(["Input", "Value"], [
    ("Member / section", "Tension; 2 ISA 55×55×10"),
    ("Factored force Tu", "375 kN"),
    ("Material strengths", "fy = 250 MPa; fu = 410 MPa"),
    ("Bolts", "5 in one line; nominal diameter 20 mm; hole 22 mm"),
    ("Connection geometry", "pitch 50 mm; end distance 33 mm; Lt = 37.8 mm"),
    ("Catalogue area", "1002 mm² per angle; 2004 mm² total"),
], [2.0, 4.8])
table(["Limit state", "App result as shown", "Status at Tu = 375 kN"], [
    ("Yielding", "455.45 kN", "Satisfies"),
    ("Rupture", "495.00 kN", "Satisfies"),
    ("Block shear", "628.58 kN", "Satisfies"),
], [1.7, 1.8, 3.3])
para("The minimum displayed resistance is 455.45 kN, governed by yielding. The displayed utilization is 375/455.45 = 0.823, so the application reports PASS for the three included checks.")
figure(SCREEN / "Screenshot 2026-10-14 183544.png", "Figure 2. Double-angle result screen captured during the activity.", 6.45)

doc.add_page_break()
doc.add_heading("4  Example one comparison with hand work", level=1)
para("The handwritten pages use the same section and load. The yielding calculation is directly comparable: 2 × 1002 × 250 / 1.10 / 1000 = 455.45 kN. The photographed rupture page records approximately 465.3 kN, while the application screenshot records 495.00 kN. The block-shear page includes overwritten values; the reading is not sufficiently clear to claim a confirmed independent reference value.")
callout("Validation note: The yielding result agrees to rounding. Rupture and block shear should not be described as fully validated by this example. Differences may arise from area deductions, β inputs, geometric interpretation, or arithmetic. Verify the definitions and substitutions against IS 800:2007 and the instructor’s accepted solution before treating either result as authoritative.")
figure(PHOTOS / "WhatsApp Image 2026-10-14 at 6.58.56 PM (1).jpeg", "Figure 3. Handwritten yielding and rupture calculations for the double-angle example.", 3.7)
doc.add_page_break()
doc.add_heading("Hand calculation evidence for example one", level=1)
figure(PHOTOS / "WhatsApp Image 2026-10-14 at 6.58.56 PM (2).jpeg", "Figure 4. Handwritten block-shear calculations; some final values are overwritten.", 4.6)

# Page 6: case B
doc.add_page_break()
doc.add_heading("5  Example two with a single angle", level=1)
para("The second submitted example uses one unequal angle connected through its 60 mm leg and carries a factored force of 200 kN. The screenshot shows the initial application output; after that capture, the β shear-lag width was corrected to include the inferred gauge for either connected-leg selection.")
table(["Input", "Value"], [
    ("Section", "ISA 90×60×6; single angle"),
    ("Factored force Tu", "200 kN"),
    ("Material strengths", "fy = 250 MPa; fu = 410 MPa"),
    ("Bolts", "5; nominal diameter 16 mm; hole 18 mm"),
    ("Connection geometry", "pitch 40 mm; end distance 27 mm; shorter leg connected; Lt = 46.1 mm"),
    ("Catalogue area", "865 mm²"),
], [2.0, 4.8])
table(["Limit state", "Notebook value", "App screenshot / updated engine"], [
    ("Yielding", "196.59 kN", "196.59 kN"),
    ("Rupture", "184.747 kN", "194.09 kN captured; 190.02 kN after β geometry correction"),
    ("Block shear", "171.259 kN", "171.26 kN"),
], [1.3, 1.7, 3.8])
para("The notebook computes β using bs = 90 + 13.9 − 6 = 97.9 mm; the app’s earlier screenshot used 90 mm. Correcting this geometric term changes β from about 1.009 to about 0.975. The updated calculation still differs from the notebook rupture strength because the net connected area in the notebook appears to be 234 mm², while the current app formula uses (60 − 18) × 6 = 252 mm². This area-definition difference remains unresolved and must be checked against the standard and class reference solution.")
doc.add_page_break()
doc.add_heading("Application result evidence for example two", level=1)
para("This result screen was captured before the β geometry correction described above. It is retained here as a record of the original activity run; the corrected post-review rupture result is shown in the comparison table on the preceding page.")
figure(SCREEN / "Screenshot 2026-10-14 185123.png", "Figure 5. Single-angle results captured before the β geometry correction.", 6.45)

doc.add_page_break()
doc.add_heading("Hand calculation evidence for example two", level=1)
para("The handwritten solution reports 196.59 kN for yielding, 184.747 kN for rupture, and 171.259 kN for block shear. The app now gives 196.59 kN, 190.02 kN, and 171.26 kN, respectively. Yielding and block shear agree to rounding. Rupture remains different because the notebook appears to use Anc = 234 mm², while the app uses Anc = 252 mm² from the entered 18 mm hole and 6 mm thickness. This is an open validation item.")
figure(PHOTOS / "WhatsApp Image 2026-10-14 at 6.56.02 PM.jpeg", "Figure 6. Handwritten single-angle yielding and rupture calculation.", 3.7)
doc.add_page_break()
doc.add_heading("Block shear hand calculation evidence", level=1)
figure(PHOTOS / "WhatsApp Image 2026-10-14 at 6.56.02 PM (1).jpeg", "Figure 7. Handwritten single-angle block-shear calculation.", 4.2)

doc.add_page_break()
doc.add_heading("6  Testing, limitations, and conclusion", level=1)
doc.add_heading("Validation and software tests", level=2)
para("Six automated tests were run with pytest and passed. They check section lookup and gross yielding substitution, double-angle catalogue handling and strength scaling, net rupture substitution, block-shear substitution, the corrected shorter-leg β geometry, and invalid/unlisted input handling. The example comparisons above are separate engineering cross-checks; the unresolved rupture differences mean the whole calculation model is not yet independently validated.")
table(["Test case", "Expected / reference", "Observed"], [
    ("Double ISA 55×55×10 gross yielding", "455.45 kN by direct substitution", "455.45 kN; agrees"),
    ("Single ISA 90×60×6 block shear", "171.259 kN from notebook", "171.26 kN; agrees to rounding"),
    ("Unsupported designation", "Reject and show ‘please select correct section’", "Handled by catalogue validation test"),
], [2.6, 2.45, 2.45])
doc.add_heading("Limitations", level=2)
para("This application is a student learning prototype, not approved structural design software. Compression selection only confirms section lookup; compression buckling checks are not implemented. It does not check bolt capacity, bearing, tear-out, gusset strength, eccentricity, connection detailing, serviceability, fatigue, load combinations, or all code requirements. Its angle table is based on the supplied 1964 steel book and may not represent current products. The actual IS 800:2007 and amendments must be consulted to verify equations, definitions, factors, material data, detailing rules, and applicability.")
doc.add_heading("Conclusion", level=2)
para("A working Python and Streamlit application was developed to guide a user through the selection and preliminary checking of a steel angle tension member. It reports yielding, rupture, and block shear separately, identifies the governing strength, and explains the result using deterministic calculations. The example checks show agreement for gross yielding and block shear in the readable comparisons, while rupture differences remain. Therefore, the activity demonstrates how a calculation workflow can be automated, but its engineering outputs require clause-by-clause verification and resolution of the noted area and geometry differences before the app can be treated as a validated calculator.")
doc.add_heading("How to run the project", level=2)
para("Open PowerShell in the project folder and run:  python -m pip install -r requirements.txt  then  python -m streamlit run app.py. The app opens locally in a browser. Enter tensile member, force and strengths, connection details, and section last. The report examples can be demonstrated with ‘2 ISA 55x55x10’ and ‘ISA 90x60x6’.")
doc.add_heading("References", level=2)
para("Bureau of Indian Standards, IS 800:2007, General Construction in Steel — Code of Practice. Verify against the licensed standard and amendments.")
para("Bureau of Indian Standards, IS SP 6 (Part 1):1964, Handbook for Structural Engineers — Steel Tables. Supplied section tables were used for educational catalogue data.")
para("Student handwritten calculations and application screenshots supplied for this FA activity; treated as comparison and demonstration evidence.")

doc.save(OUT)
print(OUT)
